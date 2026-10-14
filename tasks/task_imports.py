"""Bounded Excel parsing and additive, retry-safe task registration."""
from datetime import date, datetime
from io import BytesIO
from pathlib import PurePosixPath
import re
import uuid
from zipfile import ZipFile

from defusedxml.ElementTree import fromstring
from django.db import transaction
from django.db.models import Max
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Category, Task, TaskStatusHistory


MAX_FILE_BYTES = 2_000_000
MAX_ROWS = 500
HEADERS = ["날짜", "카테고리", "해야 할 일", "상세 메모", "상태", "중요 여부"]
EXTRA_HEADERS = {"요일", "생성일시", "수정일시"}
REQUIRED_HEADERS = {"날짜", "카테고리", "해야 할 일"}


class TaskImportError(ValueError):
    pass


def template_workbook(categories):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "일정 등록"
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="286A56")
        cell.font = Font(bold=True, color="FFFFFF")
    for column, width in zip("ABCDEF", [16, 22, 40, 52, 12, 14]):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2, max_row=MAX_ROWS + 1, max_col=6):
        row[0].number_format = "yyyy-mm-dd"
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, options in [("E", '"미정,O,X,△"'), ("F", '"일반,중요"')]:
        validation = DataValidation(type="list", formula1=options, allow_blank=True)
        validation.showErrorMessage = True
        validation.errorTitle = "입력값 확인"
        validation.error = "목록에서 값을 선택해 주세요."
        sheet.add_data_validation(validation)
        validation.add(f"{column}2:{column}{MAX_ROWS + 1}")
    names = list(categories.values_list("name", flat=True))
    guide = workbook.create_sheet("작성 안내")
    instructions = [
        ["일정 엑셀 등록 양식", "내용"],
        ["등록", "첫 번째 '일정 등록' 시트의 제목 아래에 입력하세요. 예시 행은 없습니다."],
        ["날짜", "2026-10-10 또는 엑셀 날짜. 비우면 날짜 미정 보관함에 등록합니다."],
        ["카테고리", "아래 사용 가능한 카테고리 이름을 그대로 입력하세요. 없으면 사이트에서 먼저 추가하세요."],
        ["해야 할 일", "필수, 최대 200자"],
        ["상세 메모", "선택, 최대 2,000자"],
        ["상태", "미정 / O / X / △. 비우면 미정."],
        ["중요 여부", "일반 / 중요. 비우면 일반."],
        ["파일", "xlsx만 지원합니다. 한 파일 2MB 이하, 최대 500건. 수식은 값으로 바꿔 입력하세요."],
        ["중복", "같은 날짜·카테고리·제목은 기본적으로 건너뜁니다. 미리보기에서 추가 등록을 선택할 수 있습니다."],
        ["기존 파일", "사이트에서 내려받은 일정 Excel도 사용할 수 있습니다. 요일·생성일시·수정일시는 자동으로 처리합니다."],
        ["기존 일정", "새 일정을 추가하는 기능입니다. 기존 일정, 일기, 사진은 덮어쓰지 않습니다."],
        [],
        ["사용 가능한 카테고리"],
    ]
    for row in instructions:
        guide.append(row)
    for name in names:
        guide.append([name])
        # Category names are user text, even when starting with '='.
        guide.cell(guide.max_row, 1).data_type = "s"
    guide.column_dimensions["A"].width = 26
    guide.column_dimensions["B"].width = 105
    for row in guide:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if names:
        start = len(instructions) + 1
        workbook.defined_names.add(DefinedName("ImportCategories", attr_text=f"'작성 안내'!$A${start}:$A${start + len(names) - 1}"))
        validation = DataValidation(type="list", formula1="ImportCategories", allow_blank=False)
        validation.showErrorMessage = True
        validation.error = "사용 가능한 카테고리에서 선택해 주세요."
        sheet.add_data_validation(validation)
        validation.add(f"B2:B{MAX_ROWS + 1}")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _inspect_archive(content):
    try:
        with ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > 200 or sum(item.file_size for item in members) > 20_000_000:
                raise TaskImportError("엑셀 내부 데이터가 너무 큽니다. 일정만 새 양식에 복사해 주세요.")
            names = set()
            for item in members:
                path = PurePosixPath(item.filename)
                if item.filename in names or path.is_absolute() or ".." in path.parts or item.flag_bits & 1:
                    raise TaskImportError("지원하지 않는 엑셀 파일입니다. 새 양식에 값만 복사해 주세요.")
                names.add(item.filename)
                lowered = item.filename.lower()
                if "vbaproject" in lowered or lowered.startswith("xl/externallinks/"):
                    raise TaskImportError("매크로나 외부 통합 문서 연결은 지원하지 않습니다. 값만 복사해 주세요.")
                if lowered.endswith((".xml", ".rels")):
                    # Limits also apply to shared strings/styles, before openpyxl reads them.
                    fromstring(archive.read(item), forbid_dtd=True)
            if "xl/workbook.xml" not in names:
                raise TaskImportError("올바른 xlsx 파일이 아닙니다.")
    except TaskImportError:
        raise
    except Exception as error:
        raise TaskImportError("엑셀 파일을 읽을 수 없습니다. 암호 없이 xlsx 형식으로 다시 저장해 주세요.") from error


def _text(value):
    return "" if value is None else str(value).strip()


def _date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    value = _text(value)
    if not value:
        return None
    if not re.fullmatch(r"\d{4}[-/.]\d{1,2}[-/.]\d{1,2}", value):
        raise ValueError
    year, month, day = map(int, re.split(r"[-/.]", value))
    return date(year, month, day)


def parse_workbook(content, categories):
    if not content or len(content) > MAX_FILE_BYTES:
        raise TaskImportError("xlsx 파일은 2MB 이하여야 합니다.")
    _inspect_archive(content)
    workbook = None
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        sheet = workbook.worksheets[0]
        # Do not trust the dimensions supplied by a third-party spreadsheet.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows(max_col=10)
        header_cells = next(iterator, ())
        headers = [_text(cell.value) for cell in header_cells]
        while headers and not headers[-1]:
            headers.pop()
        if any(cell.data_type in {"f", "e"} for cell in header_cells):
            raise TaskImportError("열 제목에는 수식이나 오류 값을 사용할 수 없습니다.")
        if not REQUIRED_HEADERS.issubset(headers):
            raise TaskImportError("첫 번째 시트의 첫 행에 날짜, 카테고리, 해야 할 일 열이 필요합니다. 등록 양식을 사용해 주세요.")
        if len(headers) > 9 or len(set(headers)) != len(headers) or any(name not in set(HEADERS) | EXTRA_HEADERS for name in headers):
            raise TaskImportError("열 제목이 양식과 다릅니다. 중복 열이나 불필요한 열을 지워 주세요.")
        category_map = {category.name: category for category in categories}
        status_map = {label: value for value, label in Task.Status.choices}
        results = []
        for row_number, cells in enumerate(iterator, start=2):
            if row_number > MAX_ROWS + 1001:
                raise TaskImportError("빈 행이 너무 많습니다. 입력한 일정만 새 양식에 복사해 주세요.")
            if not any(cell.value is not None and _text(cell.value) for cell in cells):
                continue
            if len(results) == MAX_ROWS:
                raise TaskImportError("한 번에 최대 500건까지 등록할 수 있습니다. 파일을 나눠 주세요.")
            values = {name: cells[index].value if index < len(cells) else None for index, name in enumerate(headers)}
            errors = []
            if any(cell.data_type in {"f", "e"} for cell in cells):
                errors.append("수식·엑셀 오류는 값으로 바꿔 주세요.")
            if any(cell.value is not None for cell in cells[len(headers):]):
                errors.append("열 제목이 없는 곳의 값을 지워 주세요.")
            try:
                task_date = _date(values.get("날짜"))
            except (ValueError, TypeError, OverflowError):
                task_date = None
                errors.append("날짜를 YYYY-MM-DD 또는 엑셀 날짜로 입력해 주세요.")
            category_name = _text(values.get("카테고리"))
            category = category_map.get(category_name)
            if not category:
                errors.append("사용 가능한 카테고리 이름을 입력하거나 사이트에서 먼저 추가해 주세요.")
            title = _text(values.get("해야 할 일"))
            details = _text(values.get("상세 메모"))
            if not title or len(title) > 200:
                errors.append("해야 할 일은 1~200자로 입력해 주세요.")
            if len(details) > 2000:
                errors.append("상세 메모는 2,000자까지 입력할 수 있습니다.")
            status_label = _text(values.get("상태")) or "미정"
            status = status_map.get(status_label.upper() if status_label.lower() in {"o", "x"} else status_label)
            if status is None:
                errors.append("상태는 미정, O, X, △ 중 하나로 입력해 주세요.")
            importance = _text(values.get("중요 여부")) or "일반"
            if importance not in {"일반", "중요"}:
                errors.append("중요 여부는 일반 또는 중요로 입력해 주세요.")
            # Invalid long cells are truncated only for display, and are never staged.
            results.append({
                "row_number": row_number, "task_date": task_date.isoformat() if task_date else "",
                "date_display": task_date.isoformat() if task_date else (_text(values.get("날짜"))[:40] or "날짜 미정"),
                "category_id": str(category.pk) if category else "", "category_name": category_name[:60],
                "title": title[:200], "details": details[:2000], "status": status,
                "status_label": status_label[:30], "is_important": importance == "중요", "errors": errors,
            })
        if not results:
            raise TaskImportError("등록할 일정이 없습니다. 열 제목 아래에 일정을 입력해 주세요.")
        return results
    except TaskImportError:
        raise
    except Exception as error:
        raise TaskImportError("엑셀 내용을 읽을 수 없습니다. 새 양식에 값만 복사해 주세요.") from error
    finally:
        if workbook is not None:
            workbook.close()


def mark_duplicates(rows):
    category_ids = {row["category_id"] for row in rows if row["category_id"]}
    titles = {row["title"] for row in rows}
    existing = {
        (day.isoformat() if day else "", str(category_id), title.strip())
        for day, category_id, title in Task.objects.filter(
            category_id__in=category_ids, title__in=titles, deleted_at__isnull=True,
        ).values_list("task_date", "category_id", "title")
    }
    seen = set()
    for row in rows:
        key = (row["task_date"], row["category_id"], row["title"])
        row["duplicate"] = "기존 일정과 중복" if key in existing else ("파일 안에서 중복" if key in seen else "")
        if not row["errors"]:
            seen.add(key)
    return rows


@transaction.atomic
def register_rows(rows, token, include_duplicates=False):
    # Consistent category lock order serializes overlapping import requests on PostgreSQL.
    category_ids = {row["category_id"] for row in rows}
    categories = {str(item.pk): item for item in Category.objects.select_for_update().filter(pk__in=category_ids).order_by("pk")}
    if any(row["category_id"] not in categories or categories[row["category_id"]].is_archived or categories[row["category_id"]].name != row["category_name"] for row in rows):
        raise TaskImportError("카테고리가 변경되었거나 보관 처리되었습니다. 파일을 다시 올려 확인해 주세요.")
    namespace = uuid.UUID(token)
    ids = [uuid.uuid5(namespace, str(row["row_number"])) for row in rows]
    if Task.objects.filter(pk__in=ids).exists():
        return 0, 0, True
    mark_duplicates(rows)
    tasks = []
    next_orders = {
        (item["task_date"], str(item["category_id"])): item["last"] + 1
        for item in Task.objects.filter(category_id__in=category_ids, deleted_at__isnull=True)
        .values("task_date", "category_id").annotate(last=Max("sort_order"))
    }
    skipped = 0
    for task_id, row in zip(ids, rows):
        if row["duplicate"] and not include_duplicates:
            skipped += 1
            continue
        day = date.fromisoformat(row["task_date"]) if row["task_date"] else None
        key = (day, row["category_id"])
        next_orders.setdefault(key, 0)
        tasks.append(Task(id=task_id, task_date=day, category_id=row["category_id"], title=row["title"], details=row["details"], status=row["status"], is_important=row["is_important"], sort_order=next_orders[key]))
        next_orders[key] += 1
    Task.objects.bulk_create(tasks, batch_size=100)
    TaskStatusHistory.objects.bulk_create([TaskStatusHistory(task=item, from_status=None, to_status=item.status) for item in tasks], batch_size=100)
    return len(tasks), skipped, False
