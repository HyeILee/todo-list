import secrets as _review_secrets
from urllib.parse import quote as _review_quote
_REVIEW_PASSWORD_1 = _review_secrets.token_urlsafe(32) + "@?"
_REVIEW_PASSWORD_2 = _review_secrets.token_urlsafe(32) + "@?"
_REVIEW_PASSWORD_3 = _review_secrets.token_urlsafe(32) + "@?"
_REVIEW_PASSWORD_4 = _review_secrets.token_urlsafe(32) + "@?"

from datetime import date, timedelta
import hashlib
import json
import sqlite3
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path
from unittest import mock
from django.contrib.auth import get_user_model
from django.contrib.sessions.models import Session
from django.core.management import CommandError, call_command
from django.test import TestCase, TransactionTestCase, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connections
from django.utils import timezone
from django.urls import reverse
from openpyxl import load_workbook
from .models import BackupLog, Category, DailyRecord, Task, TaskStatusHistory
from .forms import CategoryForm, DailyRecordForm, RecordFilterForm, TaskFilterForm, TaskForm
from .views import week_start

class AccessibleFormTests(TestCase):

    def test_forms_construct_and_reference_error_ids(self):
        for form_class in (DailyRecordForm, TaskForm, CategoryForm, TaskFilterForm, RecordFilterForm):
            for options in ({}, {'prefix': 'entry', 'auto_id': 'custom_%s'}):
                with self.subTest(form=form_class.__name__, options=options):
                    form = form_class(**options)
                    for field in form:
                        self.assertEqual(field.field.widget.attrs['aria-describedby'], f'{field.auto_id}_errors')

    def test_existing_description_is_preserved(self):
        from django import forms
        from .forms import AccessibleFormMixin

        class DescribedForm(AccessibleFormMixin, forms.Form):
            title = forms.CharField(widget=forms.TextInput(attrs={'aria-describedby': 'title-help'}))
        form = DescribedForm()
        self.assertEqual(form.fields['title'].widget.attrs['aria-describedby'], 'title-help id_title_errors')

class DashboardRegressionTests(TestCase):

    def setUp(self):
        self.user = get_user_model().objects.create_user(username='dashboard-owner', password=_REVIEW_PASSWORD_2)
        self.client.force_login(self.user)

    def assert_dashboard_renders(self):
        response = self.client.get('/')
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'dashboard.html')
        self.assertContains(response, 'for="id_today_story"')
        self.assertContains(response, 'id="id_today_story"')
        self.assertContains(response, 'aria-describedby="id_today_story_errors"')
        self.assertContains(response, 'id="id_today_story_errors"')
        return response

    def test_dashboard_without_related_data(self):
        self.assertFalse(Task.objects.exists())
        self.assertFalse(DailyRecord.objects.exists())
        response = self.assert_dashboard_renders()
        self.assertEqual(response.context['grouped_tasks'], [])
        self.assertIsNone(response.context['record'])

    def test_dashboard_with_tasks_and_record(self):
        today = timezone.localdate()
        category = Category.objects.create(name='Dashboard category')
        task = Task.objects.create(task_date=today, category=category, title='Dashboard task', status=Task.Status.DONE)
        record = DailyRecord.objects.create(record_date=today, today_story='Dashboard record')
        response = self.assert_dashboard_renders()
        self.assertContains(response, task.title)
        self.assertContains(response, record.today_story)
        self.assertEqual(response.context['record'], record)
        self.assertEqual(response.context['record_form'].instance, record)

class IsolatedWriteTests(TestCase):

    def setUp(self):
        super().setUp()
        backup = mock.patch('tasks.backups.ensure_daily_backup_before_change', return_value=True)
        self.automatic_backup = backup.start()
        self.addCleanup(backup.stop)
        directory = tempfile.TemporaryDirectory(prefix='dugyeon-test-backup-')
        self.addCleanup(directory.cleanup)
        self.enterContext(override_settings(BACKUP_DIR=Path(directory.name)))

class ChecklistTests(IsolatedWriteTests):

    def setUp(self):
        super().setUp()
        self.user = get_user_model().objects.create_user(username='owner', password=_REVIEW_PASSWORD_2)
        self.category = Category.objects.create(name='노무')

    def test_week_starts_monday_across_month_boundary(self):
        self.assertEqual(week_start(date(2026, 9, 1)), date(2026, 8, 31))
        self.assertEqual(week_start(date(2027, 1, 1)), date(2026, 12, 28))

    def test_dashboard_requires_login(self):
        response = self.client.get(reverse('home'))
        self.assertRedirects(response, f"{reverse('login')}?next=/")

    def test_security_headers_are_present(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse('home'))
        self.assertIn("default-src 'self'", response['Content-Security-Policy'])
        self.assertIn("frame-ancestors 'none'", response['Content-Security-Policy'])
        self.assertEqual(response['Referrer-Policy'], 'same-origin')
        self.assertEqual(response['X-Content-Type-Options'], 'nosniff')

    def test_error_responses_include_trace_id(self):
        response = self.client.get('/day/not-a-date/')
        self.assertEqual(response.status_code, 404)
        self.assertTrue(response['X-Error-ID'])

    @override_settings(AXES_FAILURE_LIMIT=2)
    def test_repeated_login_failures_lock_the_account(self):
        for _ in range(2):
            self.client.post(reverse('login'), {'username': self.user.username, 'password': _REVIEW_PASSWORD_3})
        response = self.client.post(reverse('login'), {'username': self.user.username, 'password': _REVIEW_PASSWORD_2})
        self.assertEqual(response.status_code, 429)
        self.assertNotIn('_auth_user_id', self.client.session)

    def test_write_without_csrf_is_rejected(self):
        from django.test import Client
        csrf_client = Client(enforce_csrf_checks=True)
        csrf_client.force_login(self.user)
        response = csrf_client.post(reverse('task_create'), {'title': '차단되어야 함', 'category': self.category.pk})
        self.assertEqual(response.status_code, 403)
        self.assertFalse(Task.objects.filter(title='차단되어야 함').exists())

    def test_invalid_day_url_is_not_silently_changed_to_today(self):
        self.client.force_login(self.user)
        self.assertEqual(self.client.get('/day/not-a-date/').status_code, 404)

    def test_development_server_rejects_public_binding(self):
        with mock.patch('django.contrib.staticfiles.management.commands.runserver.Command.handle') as start:
            with self.assertRaises(CommandError):
                call_command('runserver', '0.0.0.0:8000', use_reloader=False, verbosity=0)
            start.assert_not_called()

    def test_unapproved_mode_rejects_non_loopback_requests(self):
        response = self.client.get(reverse('home'), REMOTE_ADDR='203.0.113.10')
        self.assertEqual(response.status_code, 403)

    def test_unapproved_mode_rejects_forwarded_requests_from_local_proxy(self):
        response = self.client.get(reverse('home'), REMOTE_ADDR='127.0.0.1', HTTP_X_FORWARDED_FOR='203.0.113.10')
        self.assertEqual(response.status_code, 403)

    def test_create_task_starts_pending_and_creates_history(self):
        self.client.force_login(self.user)
        response = self.client.post(reverse('task_create'), {'task_date': '2026-09-30', 'category': self.category.pk, 'title': '자료 작성', 'details': '', 'is_important': ''})
        task = Task.objects.get()
        self.assertEqual(response.status_code, 302)
        self.assertEqual(task.status, Task.Status.PENDING)
        self.assertIsNone(task.status_history.get().from_status)

    def test_task_create_supports_dialog_fragment_and_no_script_fallback(self):
        self.client.force_login(self.user)
        dialog_response = self.client.get(reverse('task_create'), {'date': '2026-09-30', 'fragment': '1'})
        full_response = self.client.get(reverse('task_create'), {'date': '2026-09-30'})
        self.assertContains(dialog_response, 'data-task-dialog-form')
        self.assertNotContains(dialog_response, '<!doctype html>')
        self.assertContains(full_response, '<!doctype html>')

    def test_week_date_navigation_returns_fragment(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse('day', args=['2026-09-30']), HTTP_HX_REQUEST='true')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'id="day-content"')
        self.assertNotContains(response, '<!doctype html>')

    def test_assigning_inbox_task_date_preserves_creation_time(self):
        self.client.force_login(self.user)
        task = Task.objects.create(task_date=None, category=self.category, title='날짜 배정')
        created_at = task.created_at
        response = self.client.post(reverse('task_edit', args=[task.pk]), {'task_date': '2026-09-30', 'category': self.category.pk, 'title': task.title, 'details': '', 'is_important': '', 'expected_updated_at': task.updated_at.isoformat()})
        task.refresh_from_db()
        self.assertEqual(response.status_code, 302)
        self.assertEqual(task.task_date, date(2026, 9, 30))
        self.assertEqual(task.created_at, created_at)

    def test_stale_task_edit_is_rejected_and_compared(self):
        self.client.force_login(self.user)
        task = Task.objects.create(task_date=date(2026, 9, 30), category=self.category, title='초기 제목')
        Task.objects.filter(pk=task.pk).update(updated_at=timezone.now() - timedelta(seconds=1))
        task.refresh_from_db()
        old_version = task.updated_at.isoformat()
        task.title = '다른 창의 최신 제목'
        task.save()
        response = self.client.post(reverse('task_edit', args=[task.pk]), {'task_date': '2026-09-30', 'category': self.category.pk, 'title': '현재 창의 입력', 'details': '', 'is_important': '', 'expected_updated_at': old_version})
        task.refresh_from_db()
        self.assertEqual(response.status_code, 409)
        self.assertContains(response, '다른 창에서 이 업무가 수정되었습니다.', status_code=409)
        self.assertContains(response, '현재 창의 입력', status_code=409)
        self.assertEqual(task.title, '다른 창의 최신 제목')

    def test_quick_add_saves_only_title_and_category_for_selected_day(self):
        self.client.force_login(self.user)
        response = self.client.post(reverse('task_create'), {'task_date': '2026-09-30', 'category': self.category.pk, 'title': '휴대폰에서 빠른 입력'})
        task = Task.objects.get()
        self.assertEqual(response.status_code, 302)
        self.assertEqual(task.task_date, date(2026, 9, 30))
        self.assertEqual(task.title, '휴대폰에서 빠른 입력')
        self.assertEqual(task.details, '')
        self.assertFalse(task.is_important)

    def test_status_change_records_history(self):
        self.client.force_login(self.user)
        task = Task.objects.create(task_date=date(2026, 9, 30), category=self.category, title='회의자료')
        response = self.client.post(reverse('task_status', args=[task.pk]), {'status': Task.Status.DONE}, follow=True)
        task.refresh_from_db()
        self.assertRedirects(response, reverse('day', args=[task.task_date]))
        self.assertEqual(task.status, Task.Status.DONE)
        self.assertEqual((task.status_history.get().from_status, task.status_history.get().to_status), (Task.Status.PENDING, Task.Status.DONE))
        self.assertContains(response, '상태를 저장했습니다.')

    def test_status_fragment_saves_without_redirect(self):
        self.client.force_login(self.user)
        task = Task.objects.create(task_date=date(2026, 9, 30), category=self.category, title='상태 조각')
        response = self.client.post(reverse('task_status', args=[task.pk]), {'status': Task.Status.PARTIAL}, HTTP_HX_REQUEST='true')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {'status': Task.Status.PARTIAL, 'label': '△', 'task_ids': [str(task.pk)]})

    def test_inbox_is_excluded_from_week(self):
        self.client.force_login(self.user)
        Task.objects.create(task_date=None, category=self.category, title='날짜 미정')
        response = self.client.get(reverse('day', args=['2026-09-30']))
        self.assertContains(response, '0건')

    def test_soft_deleted_task_is_hidden(self):
        self.client.force_login(self.user)
        task = Task.objects.create(task_date=date(2026, 9, 30), category=self.category, title='삭제 대상')
        self.client.post(reverse('task_delete', args=[task.pk]))
        response = self.client.get(reverse('day', args=['2026-09-30']))
        self.assertNotContains(response, '삭제 대상')
        task.refresh_from_db()
        self.assertIsNotNone(task.deleted_at)

    def test_deleted_task_can_be_restored_within_30_days(self):
        self.client.force_login(self.user)
        task = Task.objects.create(task_date=date(2026, 9, 30), category=self.category, title='복구 대상')
        task.deleted_at = timezone.now()
        task.save(update_fields=['deleted_at'])
        response = self.client.post(reverse('task_restore', args=[task.pk]))
        task.refresh_from_db()
        self.assertEqual(response.status_code, 302)
        self.assertIsNone(task.deleted_at)

    def test_expired_deleted_task_cannot_be_restored(self):
        self.client.force_login(self.user)
        task = Task.objects.create(task_date=date(2026, 9, 30), category=self.category, title='만료 업무')
        task.deleted_at = timezone.now() - timedelta(days=31)
        task.save(update_fields=['deleted_at'])
        response = self.client.post(reverse('task_restore', args=[task.pk]))
        self.assertEqual(response.status_code, 404)

    def test_drag_reorder_updates_only_the_submitted_group(self):
        self.client.force_login(self.user)
        first = Task.objects.create(task_date=date(2026, 9, 30), category=self.category, title='첫 업무', sort_order=0)
        second = Task.objects.create(task_date=date(2026, 9, 30), category=self.category, title='둘째 업무', sort_order=1)
        response = self.client.post(reverse('task_reorder'), {'task_date': '2026-09-30', 'category_id': self.category.pk, 'task_ids': [str(second.pk), str(first.pk)]}, HTTP_HX_REQUEST='true')
        first.refresh_from_db()
        second.refresh_from_db()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {'ok': True, 'task_ids': [str(second.pk), str(first.pk)]})
        self.assertEqual((second.sort_order, first.sort_order), (0, 1))

    def test_drag_reorder_rejects_ids_from_another_date(self):
        self.client.force_login(self.user)
        task = Task.objects.create(task_date=date(2026, 9, 30), category=self.category, title='그룹 업무')
        other = Task.objects.create(task_date=date(2026, 10, 1), category=self.category, title='다른 그룹')
        response = self.client.post(reverse('task_reorder'), {'task_date': '2026-09-30', 'category_id': self.category.pk, 'task_ids': [str(task.pk), str(other.pk)]})
        task.refresh_from_db()
        other.refresh_from_db()
        self.assertEqual(response.status_code, 400)
        self.assertEqual((task.sort_order, other.sort_order), (0, 0))

class RecordAndExportTests(IsolatedWriteTests):

    def setUp(self):
        super().setUp()
        self.user = get_user_model().objects.create_user(username='owner', password=_REVIEW_PASSWORD_2)
        self.category = Category.objects.create(name='노무')
        self.client.force_login(self.user)

    def test_single_record_field_can_be_saved(self):
        response = self.client.post(reverse('record_save', args=[date(2026, 9, 30)]), {'today_story': '', 'memorable_sentence': '오래 기억할 문장', 'stayed_with_me': '', 'expected_updated_at': ''})
        record = DailyRecord.objects.get(record_date=date(2026, 9, 30))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(record.memorable_sentence, '오래 기억할 문장')
        self.assertEqual(record.today_story, '')

    def test_empty_record_requires_confirmation_before_delete(self):
        record = DailyRecord.objects.create(record_date=date(2026, 9, 30), today_story='기록')
        url = reverse('record_save', args=[record.record_date])
        response = self.client.post(url, {'today_story': '', 'memorable_sentence': '', 'stayed_with_me': '', 'expected_updated_at': record.updated_at.isoformat()})
        self.assertContains(response, '삭제할까요')
        self.assertTrue(DailyRecord.objects.filter(pk=record.pk).exists())
        record_version = response.context['record_version']
        self.client.post(url, {'confirm_delete': 'yes', 'expected_updated_at': record_version})
        self.assertFalse(DailyRecord.objects.filter(pk=record.pk).exists())

    def test_record_stale_update_is_rejected_without_overwriting(self):
        record = DailyRecord.objects.create(record_date=date(2026, 9, 30), memorable_sentence='처음 문장')
        DailyRecord.objects.filter(pk=record.pk).update(updated_at=timezone.now() - timedelta(seconds=1))
        record.refresh_from_db()
        old_version = record.updated_at.isoformat()
        record.memorable_sentence = '다른 창의 최신 문장'
        record.save()
        response = self.client.post(reverse('record_save', args=[record.record_date]), {'today_story': '', 'memorable_sentence': '현재 창의 입력', 'stayed_with_me': '', 'expected_updated_at': old_version})
        record.refresh_from_db()
        self.assertEqual(response.status_code, 409)
        self.assertContains(response, '현재 창의 입력', status_code=409)
        self.assertEqual(record.memorable_sentence, '다른 창의 최신 문장')

    def test_record_server_limit_preserves_invalid_input(self):
        response = self.client.post(reverse('record_save', args=[date(2026, 9, 30)]), {'today_story': 'x' * 2001, 'memorable_sentence': '', 'stayed_with_me': '', 'expected_updated_at': ''})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.context['form'].errors.as_data()['today_story'][0].code, 'max_length')
        self.assertContains(response, 'x' * 2001, status_code=400)
        self.assertFalse(DailyRecord.objects.exists())

    def test_database_rejects_empty_daily_record_rows(self):
        from django.db import IntegrityError, transaction
        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                DailyRecord.objects.create(record_date=date(2026, 9, 30))

    def test_daily_record_database_rejects_text_over_limit(self):
        from django.core.exceptions import ValidationError
        record = DailyRecord(record_date=date(2026, 9, 30), memorable_sentence='x' * 2001)
        with self.assertRaises(ValidationError):
            record.full_clean()

    def test_task_csv_applies_filters_and_has_utf8_bom(self):
        matching = Task.objects.create(task_date=date(2026, 9, 12), category=self.category, title='월말 보고', status=Task.Status.DONE)
        Task.objects.create(task_date=date(2026, 10, 2), category=self.category, title='다른 기간', status=Task.Status.DONE)
        Task.objects.create(task_date=date(2026, 9, 15), category=self.category, title='미완료', status=Task.Status.PENDING)
        response = self.client.get(reverse('export_tasks', args=['csv']), {'date_from': '2026-09-01', 'date_to': '2026-09-30', 'statuses': [Task.Status.DONE]})
        text = response.content.decode('utf-8-sig')
        self.assertTrue(response.content.startswith(b'\xef\xbb\xbf'))
        self.assertIn(matching.title, text)
        self.assertNotIn('다른 기간', text)
        self.assertNotIn('미완료', text)

    def test_record_csv_never_contains_task_data(self):
        DailyRecord.objects.create(record_date=date(2026, 9, 30), memorable_sentence='비 오는 날의 문장')
        Task.objects.create(task_date=date(2026, 9, 30), category=self.category, title='내부 업무 제목')
        response = self.client.get(reverse('export_records', args=['csv']))
        text = response.content.decode('utf-8-sig')
        self.assertIn('비 오는 날의 문장', text)
        self.assertNotIn('내부 업무 제목', text)

    def test_record_keyword_filter_searches_all_record_fields(self):
        DailyRecord.objects.create(record_date=date(2026, 9, 30), memorable_sentence='푸른 오후의 문장')
        DailyRecord.objects.create(record_date=date(2026, 9, 29), stayed_with_me='다른 하루의 음악')
        response = self.client.get(reverse('record_list'), {'q': '푸른 오후'})
        self.assertContains(response, '푸른 오후의 문장')
        self.assertNotContains(response, '다른 하루의 음악')

    def test_record_section_filter_shows_only_the_selected_field(self):
        DailyRecord.objects.create(
            record_date=date(2026, 9, 30),
            today_story='오늘 하루 본문',
            memorable_sentence='같은 날의 숨길 문장',
        )
        DailyRecord.objects.create(record_date=date(2026, 9, 29), memorable_sentence='문장만 있는 날')
        response = self.client.get(reverse('record_list'), {'section': 'today_story'})
        self.assertEqual(response.context['page_obj'].paginator.count, 1)
        self.assertContains(response, '오늘 하루 본문')
        self.assertNotContains(response, '같은 날의 숨길 문장')
        self.assertNotContains(response, '문장만 있는 날')

    def test_task_xlsx_has_expected_headers_and_date_cell(self):
        Task.objects.create(task_date=date(2026, 9, 30), category=self.category, title='자료 확인')
        response = self.client.get(reverse('export_tasks', args=['xlsx']))
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.cell(1, 1).value, '날짜')
        self.assertEqual(str(sheet.cell(2, 1).value)[:10], date(2026, 9, 30).isoformat())

    def test_password_change_updates_hash_and_keeps_current_session(self):
        response = self.client.post(reverse('password_change'), {'old_password': _REVIEW_PASSWORD_2, 'new_password1': 'Much-stronger-pass-684!', 'new_password2': 'Much-stronger-pass-684!'})
        self.user.refresh_from_db()
        self.assertRedirects(response, reverse('password_change_done'))
        self.assertTrue(self.user.check_password('Much-stronger-pass-684!'))
        self.assertEqual(self.client.get(reverse('home')).status_code, 200)

class FullBackupTests(TransactionTestCase):

    def setUp(self):
        super().setUp()
        directory = tempfile.TemporaryDirectory(prefix='dugyeon-full-backup-test-')
        self.addCleanup(directory.cleanup)
        self.enterContext(override_settings(BACKUP_DIR=Path(directory.name)))
        self.user = get_user_model().objects.create_user(username='backup-owner', password=_REVIEW_PASSWORD_1)
        self.category = Category.objects.create(name='채용')
        Task.objects.create(task_date=date(2026, 9, 30), category=self.category, title='백업 업무')
        DailyRecord.objects.create(record_date=date(2026, 9, 30), memorable_sentence='백업 기록')

    def test_daily_online_backup_is_idempotent_for_the_day(self):
        from .backups import ensure_daily_backup_before_change
        with tempfile.TemporaryDirectory() as temporary_directory:
            with override_settings(BACKUP_DIR=Path(temporary_directory)):
                self.assertTrue(ensure_daily_backup_before_change())
                self.assertTrue(ensure_daily_backup_before_change())
                logs = BackupLog.objects.filter(backup_type=BackupLog.BackupType.DAILY)
                self.assertEqual(logs.count(), 1)
                log = logs.get()
                self.assertEqual(log.status, BackupLog.Status.SUCCESS)
                backup_file = Path(temporary_directory) / log.file_name
                self.assertEqual(hashlib.sha256(backup_file.read_bytes()).hexdigest(), log.sha256)
                database = sqlite3.connect(backup_file)
                try:
                    self.assertEqual(database.execute('PRAGMA integrity_check').fetchone()[0], 'ok')
                finally:
                    database.close()

    def test_backup_requires_reauthentication_and_contains_verified_files(self):
        self.client.force_login(self.user)
        url = reverse('create_backup')
        rejected = self.client.post(url, {'password': _REVIEW_PASSWORD_4})
        self.assertEqual(rejected.status_code, 403)
        response = self.client.post(url, {'password': _REVIEW_PASSWORD_1})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/zip')
        log = BackupLog.objects.get(backup_type=BackupLog.BackupType.MANUAL)
        self.assertEqual(log.status, BackupLog.Status.SUCCESS)
        self.assertEqual(log.sha256, hashlib.sha256(response.content).hexdigest())
        with zipfile.ZipFile(BytesIO(response.content)) as archive:
            expected = {'schedule.sqlite3', '업무기록.xlsx', '업무기록.csv', '개인기록.xlsx', '개인기록.csv', 'manifest.json'}
            self.assertEqual(set(archive.namelist()), expected)
            manifest = json.loads(archive.read('manifest.json'))
            for filename, digest in manifest['sha256'].items():
                self.assertEqual(hashlib.sha256(archive.read(filename)).hexdigest(), digest)
            self.assertEqual(manifest['counts']['tasks'], 1)
            self.assertEqual(manifest['counts']['daily_records'], 1)
            from .backups import validate_backup_archive
            snapshot, validated_manifest = validate_backup_archive(response.content)
            self.assertEqual(validated_manifest['counts'], manifest['counts'])
            self.assertTrue(snapshot.startswith(b'SQLite format 3\x00'))
            with tempfile.TemporaryDirectory() as temporary_directory:
                database_file = Path(temporary_directory) / 'schedule.sqlite3'
                database_file.write_bytes(archive.read('schedule.sqlite3'))
                connection = sqlite3.connect(database_file)
                try:
                    self.assertEqual(connection.execute('PRAGMA integrity_check').fetchone()[0], 'ok')
                finally:
                    connection.close()

    @mock.patch('tasks.views.create_full_backup', side_effect=OSError('storage unavailable'))
    def test_failed_manual_backup_is_recorded(self, create_backup):
        self.client.force_login(self.user)
        response = self.client.post(reverse('create_backup'), {'password': _REVIEW_PASSWORD_1})
        log = BackupLog.objects.get(backup_type=BackupLog.BackupType.MANUAL)
        self.assertEqual(response.status_code, 503)
        self.assertEqual(log.status, BackupLog.Status.FAILED)
        self.assertIsNotNone(log.completed_at)
        self.assertIn('OSError', log.message)

    def test_invalid_restore_upload_does_not_change_current_data(self):
        self.client.force_login(self.user)
        task_count = Task.objects.count()
        response = self.client.post(reverse('restore_backup'), {'password': _REVIEW_PASSWORD_1, 'confirm_restore': 'yes', 'backup_file': SimpleUploadedFile('broken.zip', b'not a zip archive')})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(Task.objects.count(), task_count)

    def test_valid_restore_replaces_database_and_saves_pre_restore_zip(self):
        from .backups import create_full_backup, restore_database_archive
        Session.objects.create(session_key='restored-session', session_data='expired-auth-data', expire_date=timezone.now() + timedelta(days=1))
        archive_bytes, _ = create_full_backup()
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            database_path = directory / 'restore-target.sqlite3'
            with zipfile.ZipFile(BytesIO(archive_bytes)) as archive:
                database_path.write_bytes(archive.read('schedule.sqlite3'))
            database_settings = connections['default'].settings_dict.copy()
            database_settings['NAME'] = str(database_path)
            isolated = connections['default'].__class__(database_settings, 'default')
            with mock.patch.object(connections._connections, 'default', isolated), override_settings(BACKUP_DIR=directory / 'backups'):
                self.addCleanup(isolated.close)
                counts = restore_database_archive(archive_bytes)
                self.assertEqual(counts['tasks'], 1)
                self.assertTrue(list((directory / 'backups').glob('schedule-pre-restore-*.zip')))
                restored = sqlite3.connect(database_path)
                try:
                    self.assertEqual(restored.execute('PRAGMA integrity_check').fetchone()[0], 'ok')
                    self.assertEqual(restored.execute('SELECT COUNT(*) FROM django_session').fetchone()[0], 0)
                    self.assertEqual(restored.execute("SELECT COUNT(*) FROM tasks_backuplog WHERE backup_type='PRE_RESTORE'").fetchone()[0], 1)
                finally:
                    restored.close()
