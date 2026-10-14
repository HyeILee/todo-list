from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from io import BytesIO
import secrets
from unittest import mock
import uuid
from zipfile import ZIP_DEFLATED, ZipFile

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import DatabaseError, connections
from django.test import Client, TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from .exports import workbook_for_tasks
from .import_views import DRAFT_KEY
from .models import Category, DailyRecord, Task, TaskStatusHistory
from .task_imports import HEADERS, TaskImportError, parse_workbook, register_rows, template_workbook


def excel(rows, headers=HEADERS):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@override_settings(ALLOWED_HOSTS=["testserver", "127.0.0.1", "localhost"], AUTOMATIC_DATABASE_BACKUPS=False)
class ExcelImportTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="excel-fixture", password=secrets.token_urlsafe(24))
        self.client.force_login(self.user)
        self.category = Category.objects.create(name="개인")
        self.record = DailyRecord.objects.create(record_date=date(2026, 10, 6), today_story="기존 기록")

    def upload(self, rows=None, content=None):
        content = content if content is not None else excel(rows or [["2026-10-10", "개인", "병원 방문", "오후 2시", "O", "중요"]])
        return self.client.post(reverse("task_import"), {"workbook": SimpleUploadedFile("plans.xlsx", content)})

    def register(self, token=None, **extra):
        token = token or self.client.session[DRAFT_KEY]["token"]
        return self.client.post(reverse("task_import"), {"action": "register", "token": token, **extra})

    def test_login_owner_and_csrf_protect_all_entry_points(self):
        anonymous = Client()
        for name in ["task_import", "task_import_template"]:
            self.assertEqual(anonymous.get(reverse(name)).status_code, 302)
        protected = Client(enforce_csrf_checks=True)
        protected.force_login(self.user)
        self.assertEqual(protected.post(reverse("task_import"), {"action": "register"}).status_code, 403)
        other = get_user_model().objects.create_user(username="other-fixture")
        anonymous.force_login(other)
        self.assertEqual(anonymous.get(reverse("task_import")).status_code, 403)

    def test_template_has_current_categories_dropdowns_and_no_sample_task(self):
        response = self.client.get(reverse("task_import_template"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("no-store", response["Cache-Control"])
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual([cell.value for cell in workbook.worksheets[0][1]], HEADERS)
        self.assertFalse(any(cell.value for row in workbook.worksheets[0].iter_rows(min_row=2) for cell in row))
        self.assertIn("ImportCategories", workbook.defined_names)
        self.assertIn("개인", [row[0] for row in workbook.worksheets[1].values])
        self.assertEqual(len(workbook.worksheets[0].data_validations.dataValidation), 3)
        workbook.close()
        self.assertEqual(Task.objects.count(), 0)

    def test_preview_then_commit_preserves_records_and_initial_status_history(self):
        response = self.upload()
        self.assertContains(response, "병원 방문")
        self.assertIn("no-store", response["Cache-Control"])
        self.assertEqual(Task.objects.count(), 0)
        response = self.register(title="tampered client value")
        self.assertRedirects(response, reverse("task_list"))
        task = Task.objects.get()
        self.assertEqual((task.title, task.details, task.task_date, task.status, task.is_important), ("병원 방문", "오후 2시", date(2026, 10, 10), Task.Status.DONE, True))
        self.assertEqual(task.status_history.get().to_status, Task.Status.DONE)
        self.assertEqual(self.record.today_story, "기존 기록")
        self.assertNotIn(DRAFT_KEY, self.client.session)

    def test_blank_dates_go_to_inbox_and_optional_fields_default(self):
        self.upload([[None, "개인", "날짜 미정"], [datetime(2026, 11, 1), "개인", "엑셀 날짜"], ["2026.11.2", "개인", "문자 날짜"]],)
        self.register()
        undated = Task.objects.get(title="날짜 미정")
        self.assertIsNone(undated.task_date)
        self.assertEqual(undated.status, Task.Status.PENDING)
        self.assertFalse(undated.is_important)
        self.assertEqual(Task.objects.get(title="엑셀 날짜").task_date, date(2026, 11, 1))
        self.assertEqual(Task.objects.get(title="문자 날짜").task_date, date(2026, 11, 2))

    def test_maximum_batch_registers_all_tasks_and_histories(self):
        response = self.upload([[None, "개인", f"일정 {i}"] for i in range(500)])
        self.assertEqual(response.status_code, 200)
        self.register()
        self.assertEqual(Task.objects.count(), 500)
        self.assertEqual(TaskStatusHistory.objects.count(), 500)
        self.assertEqual(set(Task.objects.values_list("sort_order", flat=True)), set(range(500)))

    def test_existing_export_roundtrips_with_generated_columns(self):
        original = Task.objects.create(category=self.category, title="이전 일정", task_date=date(2026, 10, 11), details="메모", status=Task.Status.PARTIAL, is_important=True)
        response = self.upload(content=workbook_for_tasks([original]))
        self.assertContains(response, "기존 일정과 중복")
        self.register()
        self.assertEqual(Task.objects.count(), 1)
        original.refresh_from_db()
        self.assertEqual(original.status, Task.Status.PARTIAL)

    def test_duplicates_skip_by_default_and_do_not_overwrite(self):
        existing = Task.objects.create(category=self.category, task_date=date(2026, 10, 10), title="같은 일정", details="기존 메모", status=Task.Status.DONE, sort_order=7)
        self.upload([["2026-10-10", "개인", "같은 일정", "다른 메모"], ["2026-10-10", "개인", "새 일정"], ["2026-10-10", "개인", "새 일정"]])
        self.assertEqual(sum(bool(row["duplicate"]) for row in self.client.session[DRAFT_KEY]["rows"]), 2)
        self.register()
        self.assertEqual(Task.objects.count(), 2)
        existing.refresh_from_db()
        self.assertEqual((existing.details, existing.status, existing.sort_order), ("기존 메모", Task.Status.DONE, 7))
        self.assertEqual(Task.objects.get(title="새 일정").sort_order, 8)

    def test_explicit_duplicate_option_and_repeated_request_are_safe(self):
        self.upload([["2026-10-10", "개인", "같은 일정"], ["2026-10-10", "개인", "같은 일정"]])
        token = self.client.session[DRAFT_KEY]["token"]
        self.register(include_duplicates="yes")
        self.assertEqual(Task.objects.count(), 2)
        self.assertEqual(self.register(token, include_duplicates="yes").status_code, 400)
        self.assertEqual(Task.objects.count(), 2)

    def test_duplicates_are_rechecked_after_preview(self):
        self.upload()
        Task.objects.create(category=self.category, title="병원 방문", task_date=date(2026, 10, 10))
        self.register()
        self.assertEqual(Task.objects.count(), 1)

    def test_deleted_tasks_do_not_block_new_registration(self):
        Task.objects.create(category=self.category, title="병원 방문", task_date=date(2026, 10, 10), deleted_at=timezone.now())
        self.upload()
        self.register()
        self.assertEqual(Task.objects.filter(deleted_at__isnull=True).count(), 1)

    def test_any_bad_row_blocks_all_registration_and_identifies_excel_row(self):
        response = self.upload([["2026-10-10", "개인", "정상"], ["2026-02-30", "없는 분류", "x" * 201, "m" * 2001, "완료", "yes"]])
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.context["error_count"], 1)
        self.assertEqual(response.context["rows"][1]["row_number"], 3)
        self.assertEqual(len(response.context["rows"][1]["errors"]), 6)
        self.assertNotIn(DRAFT_KEY, self.client.session)
        self.assertEqual(Task.objects.count(), 0)

    def test_formulas_and_html_are_not_executed(self):
        response = self.upload([["2026-10-10", "개인", "=1+1"]])
        self.assertEqual(response.status_code, 400)
        self.assertNotIn(DRAFT_KEY, self.client.session)
        response = self.upload([["2026-10-10", "개인", "<script>alert('x')</script>"]])
        self.assertContains(response, "&lt;script&gt;")
        self.assertNotContains(response, "<script>alert('x')</script>")

    def test_missing_headers_duplicate_headers_empty_and_over_limit_files(self):
        contents = [excel([["개인", "제목"]], ["카테고리", "제목"]), excel([], HEADERS + ["날짜"]), excel([]), excel([[None, "개인", f"일정 {i}"] for i in range(501)]), b"not an xlsx"]
        for content in contents:
            with self.subTest(size=len(content)):
                self.assertEqual(self.upload(content=content).status_code, 400)
                self.assertNotIn(DRAFT_KEY, self.client.session)
        response = self.client.post(reverse("task_import"), {"workbook": SimpleUploadedFile("file.xls", b"data")})
        self.assertContains(response, "xlsx 파일을 선택", status_code=400)

    def test_bad_xml_external_workbook_and_expansion_limits_are_rejected(self):
        base = excel([[None, "개인", "일정"]])
        with ZipFile(BytesIO(base)) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        for extra_name, payload in [("malicious.xml", b'<!DOCTYPE x [<!ENTITY y "bad">]><x>&y;</x>'), ("xl/externalLinks/externalLink1.xml", b"<x/>"), ("huge.txt", b"x" * 20_000_001)]:
            output = BytesIO()
            with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
                for name, content in files.items():
                    archive.writestr(name, content)
                archive.writestr(extra_name, payload)
            with self.subTest(part=extra_name), self.assertRaises(TaskImportError):
                parse_workbook(output.getvalue(), [self.category])

    def test_cancel_expiry_tampering_and_other_session_cannot_commit(self):
        self.upload()
        token = self.client.session[DRAFT_KEY]["token"]
        self.assertEqual(self.register(str(uuid.uuid4())).status_code, 400)
        another = Client()
        another.force_login(self.user)
        self.assertEqual(another.post(reverse("task_import"), {"action": "register", "token": token}).status_code, 400)
        session = self.client.session
        draft = session[DRAFT_KEY]
        draft["created_at"] -= 1801
        session[DRAFT_KEY] = draft
        session.save()
        self.assertEqual(self.register(token).status_code, 400)
        self.upload()
        token = self.client.session[DRAFT_KEY]["token"]
        self.client.get(reverse("task_import"))
        self.assertEqual(self.register(token).status_code, 400)
        self.assertEqual(Task.objects.count(), 0)

    def test_category_change_after_preview_prevents_registration(self):
        self.upload()
        self.category.is_archived = True
        self.category.save()
        self.assertEqual(self.register().status_code, 400)
        self.assertEqual(Task.objects.count(), 0)

    def test_history_failure_rolls_back_all_new_tasks(self):
        self.upload()
        with mock.patch("tasks.task_imports.TaskStatusHistory.objects.bulk_create", side_effect=DatabaseError("fixture failure")):
            response = self.register()
        self.assertEqual(response.status_code, 503)
        self.assertEqual(Task.objects.count(), 0)
        self.assertEqual(TaskStatusHistory.objects.count(), 0)
        self.assertEqual(DailyRecord.objects.count(), 1)


class ExcelRetryTransactionTests(TransactionTestCase):
    def test_same_preview_registers_once_even_with_duplicate_option(self):
        category = Category.objects.create(name="개인")
        rows = parse_workbook(excel([["2026-10-10", "개인", "두 번 클릭"]]), [category])
        token = str(uuid.uuid4())

        def execute():
            try:
                return register_rows([dict(row) for row in rows], token, include_duplicates=True)
            finally:
                connections.close_all()

        if connections["default"].vendor == "postgresql":
            with ThreadPoolExecutor(max_workers=2) as pool:
                results = list(pool.map(lambda _: execute(), range(2)))
        else:
            results = [register_rows(rows, token, True), register_rows(rows, token, True)]
        self.assertEqual(sum(result[0] for result in results), 1)
        self.assertEqual(sum(result[2] for result in results), 1)
        self.assertEqual(Task.objects.count(), 1)
        self.assertEqual(TaskStatusHistory.objects.count(), 1)
