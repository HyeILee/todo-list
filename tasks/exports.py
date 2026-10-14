import csv
from datetime import datetime
from io import BytesIO, StringIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Task


WEEKDAYS = ["월", "화", "수", "목", "금", "토", "일"]


def _local_datetime(value):
    return timezone.localtime(value).replace(tzinfo=None) if value else ""


def task_export_rows(tasks):
    rows = []
    for task in tasks:
        rows.append(
            [
                task.task_date,
                WEEKDAYS[task.task_date.weekday()] if task.task_date else "",
                task.category.name,
                task.title,
                task.details,
                Task.Status(task.status).label,
                "중요" if task.is_important else "일반",
                _local_datetime(task.created_at),
                _local_datetime(task.updated_at),
            ]
        )
    return rows


def record_export_rows(records):
    return [
        [
            record.record_date,
            WEEKDAYS[record.record_date.weekday()],
            record.today_story,
            record.memorable_sentence,
            record.stayed_with_me,
            _local_datetime(record.updated_at),
        ]
        for record in records
    ]


def build_xlsx(sheet_name, headers, rows, widths):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    sheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="286A56")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + column)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
        if row[0].value:
            row[0].number_format = "yyyy-mm-dd"
    if sheet.max_row > 1:
        sheet.auto_filter.ref = f"A1:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_csv(headers, rows):
    output = StringIO(newline="")
    writer = csv.writer(output, lineterminator="\r\n")
    writer.writerow(headers)
    writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig")


def export_headers(kind):
    if kind == "tasks":
        return ["날짜", "요일", "카테고리", "해야 할 일", "상세 메모", "상태", "중요 여부", "생성일시", "수정일시"]
    return ["날짜", "요일", "오늘 하루는", "기억 속 문장", "유독 내 곁에", "수정일시"]


def export_filename(kind, extension):
    timestamp = timezone.localtime().strftime("%Y%m%d_%H%M")
    prefix = "업무기록" if kind == "tasks" else "개인기록"
    return f"{prefix}_{timestamp}.{extension}"


def workbook_for_tasks(tasks):
    return build_xlsx("업무 기록", export_headers("tasks"), task_export_rows(tasks), [14, 8, 16, 36, 48, 12, 12, 20, 20])


def workbook_for_records(records):
    return build_xlsx("개인 기록", export_headers("records"), record_export_rows(records), [14, 8, 60, 45, 45, 20])


def csv_for_tasks(tasks):
    return build_csv(export_headers("tasks"), task_export_rows(tasks))


def csv_for_records(records):
    return build_csv(export_headers("records"), record_export_rows(records))
